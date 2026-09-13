"""Offline tests for the session-only Candidate Profile feature."""

import io
import unittest
from unittest.mock import MagicMock

import pandas as pd
from docx import Document

from src.candidate_profile import (
    CandidateUploadError,
    build_candidate_profile,
    clear_candidate_session,
    extract_candidate_document,
    generate_application_guidance,
    match_candidate_to_posting,
    rank_candidate_matches,
    validate_candidate_upload,
)


class CandidateUploadTests(unittest.TestCase):
    def test_txt_is_extracted_and_skills_detected(self):
        result = extract_candidate_document(
            "resume.txt", b"Data student with Python, SQL, and communication experience."
        )
        self.assertEqual(result["file_type"], "txt")
        self.assertEqual(result["detected_skills"], ["Communication", "Python", "SQL"])
        self.assertNotIn("content", result)

    def test_docx_paragraphs_and_tables_are_extracted(self):
        document = Document()
        document.add_paragraph("Python data student")
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "SQL project"
        stream = io.BytesIO()
        document.save(stream)
        result = extract_candidate_document("resume.docx", stream.getvalue())
        self.assertIn("Python data student", result["text"])
        self.assertIn("SQL project", result["text"])

    def test_invalid_type_empty_binary_and_oversize_are_rejected(self):
        with self.assertRaises(CandidateUploadError):
            validate_candidate_upload("resume.exe", b"content")
        with self.assertRaises(CandidateUploadError):
            validate_candidate_upload("resume.txt", b"")
        with self.assertRaises(CandidateUploadError):
            validate_candidate_upload("resume.txt", b"a" * (5 * 1024 * 1024 + 1))
        with self.assertRaises(CandidateUploadError):
            validate_candidate_upload("resume.pdf", b"not a pdf")
        with self.assertRaises(CandidateUploadError):
            validate_candidate_upload("resume.txt", b"text\x00binary")


class CandidateMatchingTests(unittest.TestCase):
    def setUp(self):
        self.profile = build_candidate_profile(
            skills=["Python", "SQL", "ImaginaryDB"],
            education="State University",
            experience="Data project",
        )
        self.postings = pd.DataFrame([
            {"posting_id": "2", "company": "Beta", "job_title": "ML Intern",
             "skills_extracted": "Python|Machine Learning"},
            {"posting_id": "1", "company": "Alpha", "job_title": "Data Intern",
             "skills_extracted": "Python|SQL"},
        ])

    def test_reviewed_profile_is_bounded_and_catalog_backed(self):
        self.assertEqual(self.profile["skills"], ["Python", "SQL"])
        self.assertEqual(self.profile["education"], "State University")

    def test_single_posting_match_uses_deterministic_calculation(self):
        result = match_candidate_to_posting(self.profile, self.postings.iloc[0])
        self.assertEqual(result["match_percentage"], 50.0)
        self.assertEqual(result["missing_skills"], ["Machine Learning"])
        self.assertEqual(result["method"], "deterministic_catalog_skill_overlap")

    def test_ranked_matches_are_descending(self):
        result = rank_candidate_matches(self.profile, self.postings)
        self.assertEqual(result.iloc[0]["posting_id"], "1")
        self.assertEqual(result.iloc[0]["match_percentage"], 100.0)


class CandidatePrivacyAndGuidanceTests(unittest.TestCase):
    def test_clear_removes_all_candidate_values_and_resets_uploader(self):
        state = {
            "candidate_document": {"text": "private"},
            "candidate_skills": ["Python"],
            "candidate_guidance": "private guidance",
            "candidate_upload_2": b"original file bytes",
            "candidate_upload_version": 2,
            "unrelated": "keep",
        }
        clear_candidate_session(state)
        self.assertNotIn("candidate_document", state)
        self.assertNotIn("candidate_guidance", state)
        self.assertNotIn("candidate_upload_2", state)
        self.assertEqual(state["candidate_upload_version"], 3)
        self.assertEqual(state["unrelated"], "keep")

    def test_guidance_is_stateless_and_uses_reviewed_fields(self):
        client = MagicMock()
        client.responses.create.return_value.output_text = "Tailored suggestion"
        factory = MagicMock()
        factory.return_value.__enter__.return_value = client
        profile = build_candidate_profile(skills=["Python"], experience="Reviewed experience")
        posting = {
            "posting_id": "1", "company": "Example", "job_title": "Data Intern",
            "location": "Phoenix", "skills_extracted": "Python|SQL",
            "description": "Use SQL", "source_url": "https://example.com/1",
        }
        result = generate_application_guidance(
            "key", "model", profile, posting, "resume", factory
        )
        self.assertEqual(result, "Tailored suggestion")
        request = client.responses.create.call_args.kwargs
        self.assertFalse(request["store"])
        self.assertIn("Reviewed experience", request["input"])
        self.assertNotIn("key", request["input"])

    def test_invalid_guidance_type_is_rejected_before_api_call(self):
        factory = MagicMock()
        with self.assertRaises(ValueError):
            generate_application_guidance(
                "key", "model", build_candidate_profile(skills=[]), {}, "application", factory
            )
        factory.assert_not_called()


if __name__ == "__main__":
    unittest.main()
